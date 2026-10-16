"""Create the review PDF and editable PPTX using notebook outputs, without retyping chart values."""
from pathlib import Path
import json,math,re,html
import pandas as pd
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor
from PIL import Image
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
root=Path(__file__).resolve().parents[1]
summary=json.loads((root/'outputs/summary.json').read_text())
res=pd.read_csv(root/'outputs/growth_results.csv');rob=pd.read_csv(root/'outputs/robustness.csv');dates=pd.read_csv(root/'research/reform_dates.csv')
sources=json.loads((root/'research/source_register.json').read_text());index=pd.read_csv(root/'figures/chart_index.csv')
label='LDD2026_SuwarnaPyakurel';out=root/'deck';out.mkdir(exist_ok=True)
W,H=960,540
pdfmetrics.registerFont(TTFont('D','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DB','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))
pdf=canvas.Canvas(str(out/f'{label}_Deck.pdf'),pagesize=(W,H))
pdf.setTitle('Tariff cuts, uneven income paths — Suwarna Pyakurel')
pdf.setAuthor('Suwarna Pyakurel')
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
INK='#162D3D';BLUE='#176B87';GOLD='#C28A24';RED='#B33F46';GREY='#536671';PALE='#F1F5F6'
slides=[];current=None;slide_no=0;notes=[]
def rgb(c):return RGBColor.from_string(c.lstrip('#'))
def box(x,y,w,h,color):
 pdf.setFillColor(HexColor(color));pdf.rect(x,H-y-h,w,h,fill=1,stroke=0)
 sh=current.shapes.add_shape(1,Inches(x/72),Inches(y/72),Inches(w/72),Inches(h/72));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background()
def text(x,y,w,h,s,size=18,bold=False,color=INK):
 s=str(s);font='DB' if bold else 'D';escaped=html.escape(s).replace('\n','<br/>')
 chosen=size
 while True:
  style=ParagraphStyle('p',fontName=font,fontSize=chosen,leading=chosen*1.32,textColor=HexColor(color),spaceAfter=0)
  par=Paragraph(escaped,style);aw,ah=par.wrap(w,h)
  if ah<=h or chosen<=7:break
  chosen-=.5
 assert ah<=h+1,f'Text overflow slide{slide_no}: {s[:80]}'
 par.drawOn(pdf,x,H-y-ah)
 shape=current.shapes.add_textbox(Inches(x/72),Inches(y/72),Inches(w/72),Inches(h/72));tf=shape.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for j,line in enumerate(s.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(chosen);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(0)
 return chosen
def image(num,x,y,w,h):
 path=root/'figures'/f'Fig{num}.png'
 with Image.open(path) as im:iw,ih=im.size
 scale=min(w/iw,h/ih);dw,dh=iw*scale,ih*scale;xx=x+(w-dw)/2;yy=y+(h-dh)/2
 pdf.drawImage(str(path),xx,H-yy-dh,width=dw,height=dh,mask='auto')
 current.shapes.add_picture(str(path),Inches(xx/72),Inches(yy/72),width=Inches(dw/72),height=Inches(dh/72))
def start(title,kicker='LIBERTY DATA DRIVE 2026',appendix=False):
 global current,slide_no
 if current is not None:pdf.showPage()
 slide_no+=1;current=prs.slides.add_slide(prs.slide_layouts[6]);box(0,0,W,H,'#FFFFFF');box(0,0,12,H,BLUE)
 text(36,17,890,18,kicker,10,True,BLUE);text(36,46,890,66,title,27,True)
 text(36,516,740,14,'Suwarna Pyakurel · Solo · Participant review copy · Organiser approval pending',8,color=GREY)
 text(840,516,85,14,('A' if appendix else '')+str(slide_no),8,color=GREY)
 slides.append({'page':slide_no,'title':title,'appendix':appendix})
def note(s):
 notes.append({'page':slide_no,'notes':s})
 try:current.notes_slide.notes_text_frame.text=s
 except Exception:pass
# 1
start('Tariff cuts, uneven income paths')
text(48,132,850,155,'Among India, China, Mauritius, Thailand, the Philippines, Pakistan, Côte d’Ivoire and Senegal, what happened to GNI per capita in the ten years after tariff-cutting episodes in 1991–1998—and what does this suggest for Nepal’s tariff schedule?',25,True)
box(48,320,850,115,PALE)
text(66,335,810,88,'Ten countries pass the policy-data screen. Eight have documented episode years for the main comparison. Outcomes are mixed; the analysis does not identify a causal growth effect.',19)
text(48,456,850,28,'Outcome source: UNDP HDR25, GNI per capita in 2021 PPP dollars. Event outcome coverage: 1990–2008.',11,color=GREY)
note('Research question covers the eight dated cases. The screen uses raw tariff changes 1995–2005, which is different from the implementation window. All dates are annual episode conventions; finer temporal precision differs by source. No causal estimate is claimed.')
# 2
start('A policy-data rule, followed by documentary dating')
text(45,119,870,55,'Screen: mean tariff decline ≥20 percentage points, 1995–2005, with both tariff endpoints and some UNDP GNI data. All ten qualifying countries stay visible.',18)
main=dates[dates.main_sample.eq(1)]
for k,row in enumerate(main.itertuples()):
 col=k//4;rr=k%4;x=48+col*448;y=198+rr*43
 text(x,y,260,30,row.country,18,True);text(x+270,y,100,30,str(row.reform_year),18,color=BLUE)
text(48,390,860,63,'Bangladesh and Kenya: timing precision insufficient for the main tariff-event definition; retained in screening and a separate date grid. CIV/SEN share a regional reform.',15)
text(48,467,860,28,'Sources: EFW raw 4Aii; budget, executive-order, WTO, PIIE, SPDC and IMF records. Full locations: source appendix and date ledger.',10,color=GREY)
note('Dates: IND1991 budget; CHN1992 implemented cuts; MUS1994 consolidated schedule; THA1995 January programme; PHL1995 EO264 effective August28; PAK1997 tariff peak cut; CIV1998 March transition; SEN1998 April national phase. Sources differ in precision. The date-quality decisions followed earlier exploratory outcome inspection, so no preregistration claim is made. Programme start is not necessarily first liberalisation ever. BGD/KEN are not omitted because of outcomes.')
# 3
start('Large tariff declines are visible in the policy data')
image(1,35,111,890,380)
note('Figure1; notebook cell '+str(index.set_index('figure').loc['Fig1','notebook_cell'])+'. All ten screened countries, including BGD/KEN. Dataset observation years rather than presumed annual historical observations. The raw EFW mean is not a statutory maximum or a chain-linked rating.')
# 4
start('Use the required chain-linked panel as context')
image(2,35,111,890,380)
note('Figure2 analyses EFW Panel Dataset Area4 independently of the raw tariff series. Area4 combines several trade institutions. It must not be called a tariff-only measure. Exact organiser acceptance of the combined panel/raw-subcomponent interpretation is pending; the declaration asks this explicitly.')
# 5
start(f'Income rose in {summary["n_post_positive"]} of {summary["n"]} cases; that does not establish an effect')
image(4,35,110,700,385)
text(752,143,170,45,'Event alignment',17,True,color=BLUE)
text(752,194,170,105,'Every series equals 100 at reform by construction. That meeting point is not evidence.',14)
text(752,318,170,136,'Solid median changes membership before reform. Dotted median holds the five full-window countries fixed.',14)
note('Figure4. No tariff-driven income gain is identified. India, China and Mauritius lack five full pre-years in UNDP; the balanced subset therefore differs. Broader market reforms and macroeconomic conditions are plausible alternative explanations. Raw full-period income plots are in Appendix B.')
# 6
start(f'Growth accelerated in {summary["accelerated"]} of {summary["full_pre_n"]} cases with full prehistory')
image(5,35,110,890,380)
note('Figure5. Full five-year pre-CAGR compared to ten-year post-CAGR; missing pre-bars are not zero. The comparator is median non-case CAGR over each country’s same calendar years, with valid sample sizes 148–152. It is not an untreated control. China and India grew quickly but do not enter the full-prehistory acceleration count.')
# 7
start(f'Only {summary["above"]}/{summary["n"]} exceeded the broad comparator')
image(6,28,117,610,345);image(7,642,118,290,195)
text(658,331,265,105,f'Median excess: {summary["median_excess_pp"]:+.2f} pp/year.\nDating shifts change the sign of the median. The descriptive premium is fragile.',16,True)
text(44,475,875,25,'Figures 6–7. Contrary cases remain. Comparator countries can also reform; this comparison does not isolate tariff effects.',11,color=GREY)
note('Figure6 red cases are Pakistan, Thailand, Senegal and Côte d’Ivoire. Figure7 valid denominators exclude unavailable pre-1990 starts. Thailand’s Asian financial crisis and Côte d’Ivoire’s political conflict are realistic alternative explanations, documented by WTO1999 and World Bank2003. The analysis does not quantify their contributions or claim they fully explain the patterns.')
# 8
start('The conclusion survives only in a cautious form')
without_china=rob[rob['check'].eq('without China')].iloc[0];without_waemu=rob[rob['check'].eq('without both WAEMU')].iloc[0]
cards=[('Dependence on cases',f'Without China: {int(without_china.above)}/{int(without_china.n)} above comparator; median {without_china.median_excess_pp:+.2f} pp/year. Removing both WAEMU countries changes the composition, not the main sample.'),('Timing uncertainty',f'Adding Bangladesh/Kenya under all 12 stated timing combinations gives {summary["timing_grid_above_min"]}/10 above comparator. Those are assumed event years, not newly verified laws.'),('Competing explanations','China/India’s wider reforms; Thailand’s financial crisis; Côte d’Ivoire’s instability and conflict. No matched counterfactual separates these from tariff changes.'),('Transfer to Nepal','Historical national averages cannot identify an optimal Nepal tariff or predict revenue, jobs, prices or income gains. Raw means and statutory tariff peaks are different objects.')]
for k,(title,body) in enumerate(cards):
 x=45+(k%2)*448;y=123+(k//2)*179
 box(x,y,420,157,PALE);text(x+15,y+12,390,30,title,18,True,BLUE);text(x+15,y+49,390,95,body,14)
note('Sensitivity outputs include five-year post-windows, shifts, excluding China, dropping both WAEMU cases, retaining one WAEMU case, using completion2000, and threshold changes limited to documented cases. A lower numerical threshold adds countries without sourced dates; its dated subset is not a complete expanded sample. These are limitations, not hidden successes.')
# 9
start('One modification: a transparent ceiling on tariff peaks')
image(8,465,128,465,255)
text(44,119,400,145,'Nepal: amend Schedule 1 of the Customs Tariff Act 2081, as amended, to cap standard non-preferential ad valorem basic duties in HS25–97 at 20% over three budgets. Retain lower rates.',18,True)
text(44,283,400,110,'Authority: Ministry of Finance proposes; Federal Parliament enacts; Customs administers. Mechanism: reduce the tax wedge on affected imports, widening choice and competition.',15)
text(480,400,430,77,'Figures 1–2 show historical policy changes; 4–7 limit the growth claim; 8 shows Nepal’s different starting point. A 20% ceiling is proposed, not an estimated optimum.',13)
text(44,434,400,58,'Current legal source: official 2026/27 tariff compilation, pages 4–5 and 215; Customs Tariff Act 2081.',10,color=GREY)
note('Exactly one modification. Scope: non-preferential basic ad valorem duties only, HS chapters25–97; currently lower rates retained. Specific/mixed duties, other charges, VAT, excise, import prohibitions and preferences are outside the proposed amendment. The three-budget transition and20% ceiling are normative design choices; no affected-line count or fiscal/growth effect is estimated. Current compilation page215 shows higher non-agricultural basic duties, but a full product-level appraisal is still needed before implementation. This goes beyond existing band simplification; it is not described as an already implemented ceiling.')
# 10
start('More freedom to exchange, with costs stated openly')
text(48,128,850,100,'Liberty principle: purchasers and sellers should have wider freedom to exchange under predictable rules. Protection gives concentrated benefits to selected producers while spreading costs across buyers.',22,True)
box(48,257,850,130,PALE)
text(65,272,810,100,'Strongest objection: tariff revenue may fall before public finances adjust, while protected firms and workers face concentrated losses. A gradual transition does not remove these costs. Our data do not quantify them.',19)
text(48,416,850,65,'Policy judgement: lower remaining tariff peaks cautiously; do not promise a growth premium. The historical evidence is mixed and cannot establish the proposed ceiling’s welfare effect.',17,True,BLUE)
note('The normative liberty case is presented separately from the descriptive empirical evidence. The recommendation cannot be justified solely by the positive income trajectories. External submission checks remain: organiser acceptance of dataset release/interpretation, approved declaration, exact deadline, participant review and judge-access test. Those are not computational checks.')
# 11
start('Appendix A · Chart index',appendix=True)
text(45,112,860,35,'Cell numbers count markdown and code cells; all chart values are generated by the notebook.',13,color=GREY)
for k,row in enumerate(index.itertuples()):
 y=163+k*37
 text(45,y,65,30,row.figure,13,True);text(115,y,600,30,row.title,12);text(728,y,185,30,f'page {row.slide} · cell {row.notebook_cell}',12)
text(45,467,865,28,'Exact data: outputs/FigN_data.csv. Charts: figures/FigN.png and .pdf. Working file includes original dataset bytes.',11,color=GREY)
note('Source of this index is figures/chart_index.csv produced from actual notebook cell positions. Figure3 is Appendix B at page12. Sources follow in Appendix C. All eight figures have independent data exports.')
# 12
start('Appendix B · Full-period raw income histories',appendix=True)
image(3,35,112,890,380)
note('Figure3: complete supplied UNDP GNI series1990–2023, including all ten screened countries. Different labelled panel scales. Event-window conclusions in main slides use only their defined post periods. This figure provides inspection context, not an extra causal claim.')
# 13–15
for groupnum,startidx in enumerate(range(0,len(sources),7),1):
 subset=sources[startidx:startidx+7]
 start(f'Appendix C{groupnum} · Public sources and access dates',appendix=True)
 for j,s in enumerate(subset):
  y=116+j*53
  text(43,y,86,42,s['id'],9,True,BLUE)
  text(135,y,780,23,f"{s['publisher']} · {s['title']} · {s['version_year']}",9)
  text(135,y+23,780,26,f"Accessed {s['accessed']}. {s['location']}",8,color=GREY)
  pdf.linkURL(s['url'],(135,H-y-23,915,H-y),relative=0,thickness=0)
 text(45,493,870,17,'Source titles are clickable in PDF. Full URLs, source files, checksums and locations: research/sources.md and source_register.csv.',8,color=GREY)
 note('Full bibliography with literal URLs is supplied in research/sources.md and the working-file package. All sources are public; full books need not be bought because the cited China chapter is a free excerpt. Archive readability and date precision vary and are disclosed. No statistical chart values are taken from policy sources.')
# 16
start('Appendix D · AI disclosure and reproducibility',appendix=True)
text(46,121,865,120,'OpenAI Codex assisted with inspecting datasets, writing and debugging code, checking reproducibility, researching public policy sources, drafting analysis documentation and generating the presentation. The participant is responsible for reviewing the work and explaining all submitted claims.',18)
text(46,269,865,92,'The working notebook embeds both original datasets and the date ledger. Run all cells from an empty directory or fresh Colab runtime. Exact figure tables, source locations, package versions and validation reports are included.',17)
text(46,384,865,95,'Solo entry: no team contribution statement. Organiser declaration approval, dataset-scope acceptance, final deadline and signed-out sharing checks remain external pending items. No submission or organiser contact has been made by the assistant.',16,color=GREY)
note('This disclosure describes AI assistance; it does not assert that participant review has already occurred. The prepared package is a participant review copy. Fresh empty-directory execution is checked in outputs/clean_runtime_verification.json. Computational reproducibility does not imply causal identification or organiser approval.')
pdf.save();prs.save(out/f'{label}_Deck.pptx')
(out/'slide_manifest.json').write_text(json.dumps(slides,ensure_ascii=False,indent=2))
(out/'speaker_notes.md').write_text('\n\n'.join(f"## Page {n['page']}\n\n{n['notes']}" for n in notes))
assert sum(not x['appendix'] for x in slides)==10
print('Built',len(slides),'pages, 10 content slides; PDF and editable PPTX.')
